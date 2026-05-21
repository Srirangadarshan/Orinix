import os
import re


TEXT_EXTENSIONS = {
  ".md", ".txt", ".rst", ".org",
  ".py", ".sh", ".js", ".ts", ".rs", ".go", ".rb", ".php",
  ".c", ".h", ".cpp", ".hpp", ".java", ".swift", ".kt", ".scala",
  ".css", ".html", ".htm", ".xml", ".json", ".toml", ".yaml", ".yml",
  ".conf", ".ini", ".cfg", ".env", ".gitignore", ".editorconfig",
  ".csv", ".tsv", ".log", ".sql", ".r", ".lua",
}


def extract_text(filepath: str, max_size: int = 10 * 1024 * 1024) -> str | None:
  try:
    stat = os.stat(filepath)
    if stat.st_size > max_size:
      return None
    if stat.st_size == 0:
      return None
  except OSError:
    return None

  ext = os.path.splitext(filepath)[1].lower()

  if ext in TEXT_EXTENSIONS:
    return _read_text(filepath)
  if ext == ".pdf":
    return _extract_pdf(filepath)
  if ext == ".docx":
    return _extract_docx(filepath)
  if ext == ".xlsx":
    return _extract_xlsx(filepath)
  if ext == ".pptx":
    return _extract_pptx(filepath)

  return None


def _read_text(filepath: str) -> str | None:
  try:
    with open(filepath, "r", encoding="utf-8", errors="replace") as f:
      return f.read()
  except Exception:
    return None


def _extract_pdf(filepath: str) -> str | None:
  try:
    from pypdf import PdfReader
    reader = PdfReader(filepath)
    pages = []
    for page in reader.pages:
      text = page.extract_text()
      if text:
        pages.append(text)
    return "\n".join(pages) if pages else None
  except Exception:
    return None


def _extract_docx(filepath: str) -> str | None:
  try:
    from docx import Document
    doc = Document(filepath)
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paras) if paras else None
  except Exception:
    return None


def _extract_xlsx(filepath: str) -> str | None:
  try:
    from openpyxl import load_workbook
    wb = load_workbook(filepath, read_only=True, data_only=True)
    lines = []
    for sheet in wb.worksheets:
      for row in sheet.iter_rows(values_only=True):
        vals = [str(v) for v in row if v is not None]
        if vals:
          lines.append("  ".join(vals))
    return "\n".join(lines) if lines else None
  except Exception:
    return None


def _extract_pptx(filepath: str) -> str | None:
  try:
    from pptx import Presentation
    prs = Presentation(filepath)
    lines = []
    for slide in prs.slides:
      for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text.strip():
          lines.append(shape.text)
    return "\n".join(lines) if lines else None
  except Exception:
    return None


def chunk_text(text: str, max_chars: int = 2000, overlap: int = 200) -> list[str]:
  if not text.strip():
    return []

  if len(text) <= max_chars:
    return [text.strip()]

  paragraphs = re.split(r"\n\s*\n", text)
  chunks = []
  current = []

  for para in paragraphs:
    para = para.strip()
    if not para:
      continue
    current_len = sum(len(p) for p in current)
    if current_len + len(para) > max_chars and current:
      chunks.append("\n\n".join(current))
      overlap_text = current[-overlap // len(current[-1]) + 1:] if current else []
      current = list(overlap_text)
    current.append(para)

  if current:
    chunks.append("\n\n".join(current))

  return [c.strip() for c in chunks if c.strip()]
